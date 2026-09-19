#!/usr/bin/env python3
"""Generate the Labor Library catalog from the source-audited workbook. Does not invent rows."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOKS = [
    Path("/opt/cursor/project-storages/frey2535-estim8r/internal/Estim8r_Source_Audited_Electrical_Labor_Database.xlsx"),
    Path.home() / "Library/Application Support/Cursor/AgentStores/cursor_agent_stores/bc-2420d95f-c39f-426a-b160-1a7c151124fe/files/internal/Estim8r_Source_Audited_Electrical_Labor_Database.xlsx",
    ROOT / "docs/labor/Estim8r_Source_Audited_Electrical_Labor_Database.xlsx",
]
DATA_JS = ROOT / "src/domain/labor/auditedLibraryData.js"
SQL = ROOT / "supabase/migrations/20260919140000_seed_full_electrical_labor.sql"


def num(value):
    if value is None or value == "":
        return None
    return round(float(value), 4)


def sql_str(value):
    if value is None:
        return "null"
    return "'" + str(value).replace("'", "''") + "'"


def sql_num(value):
    if value is None:
        return "null"
    return repr(value)


def item_uuid(n):
    return f"00000000-0000-4000-a000-{n:012d}"


def unit_uuid(n):
    return f"00000000-0000-4000-a000-{100000 + n:012d}"


def find_workbook():
    for path in WORKBOOKS:
        if path.exists():
            return path
    raise SystemExit("Workbook not found")


def main():
    path = find_workbook()
    wb = load_workbook(path, data_only=True)
    ws = wb["Labor Database"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {name: i for i, name in enumerate(headers)}
    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        ready = str(row[idx["Production Ready"]] or "").strip().upper() == "YES"
        status = "verified" if ready else "unverified"
        source_name = row[idx["Source Type"]] or ("NECA MLU public sample" if ready else "Estim8r original model baseline")
        if ready:
            normal = num(row[idx["Verified Normal MH"]])
            difficult = num(row[idx["Verified Difficult MH"]])
            very = num(row[idx["Verified Very Difficult MH"]])
            source_type = "published_reference"
            unit_id_suffix = "neca-mlu-public-sample"
        else:
            normal = num(row[idx["Model Normal MH"]])
            difficult = num(row[idx["Model Difficult MH"]])
            very = num(row[idx["Model Very Difficult MH"]])
            source_type = "estim8r_standard"
            unit_id_suffix = "model-baseline"
        item_id = row[idx["ID"]]
        items.append({
            "id": item_id,
            "trade": "Electrical",
            "category": row[idx["Category"]] or "",
            "subcategory": row[idx["Subcategory"]] or "",
            "item_name": row[idx["Item / Installation Task"]] or "",
            "description": row[idx["Includes"]] or "",
            "material_type": row[idx["Material / System"]] or "",
            "size": row[idx["Size / Rating"]] or "",
            "unit": row[idx["Unit"]] or "EA",
            "default_crew": row[idx["Crew Type"]] or "",
            "active": str(row[idx["Active"]] or "Yes").strip().lower() == "yes",
            "labor_units": [{
                "id": f"{item_id}-{unit_id_suffix}",
                "labor_item_id": item_id,
                "source_type": source_type,
                "source_name": source_name,
                "source_year": row[idx["Source Edition"]] or "",
                "source_reference": row[idx["Verification Source URL"]] or row[idx["Source URL"]] or "",
                "normal_mh": normal,
                "difficult_mh": difficult,
                "very_difficult_mh": very,
                "verification_status": status,
                "production_allowed": ready,
                "notes": row[idx["Verification Notes"]] or "",
            }],
        })

    if len(items) != 1921:
        raise SystemExit(f"Expected 1921 labor rows, got {len(items)}")

    DATA_JS.write_text(
        "/* Generated from Estim8r_Source_Audited_Electrical_Labor_Database.xlsx Labor Database. Do not edit by hand. */\n"
        f"export const AUDITED_LABOR_ITEMS = {json.dumps(items, ensure_ascii=False)};\n",
        encoding="utf-8",
    )

    item_values = []
    unit_values = []
    for row in items:
        n = int(row["id"].split("-")[1])
        unit = row["labor_units"][0]
        item_values.append(
            f"  ({sql_str(item_uuid(n))}, {sql_str(row['trade'])}, {sql_str(row['category'])}, {sql_str(row['subcategory'])}, "
            f"{sql_str(row['item_name'])}, {sql_str(row['description'])}, {sql_str(row['material_type'])}, {sql_str(row['size'])}, "
            f"{sql_str(row['unit'])}, 1, {sql_str(row['default_crew'])}, {'true' if row['active'] else 'false'})"
        )
        unit_values.append(
            f"  ({sql_str(unit_uuid(n))}, {sql_str(item_uuid(n))}, {sql_str(unit['source_type'])}, {sql_str(unit['source_name'])}, "
            f"{sql_str(unit['source_year'])}, {sql_str(unit['source_reference'])}, {sql_num(unit['normal_mh'])}, "
            f"{sql_num(unit['difficult_mh'])}, {sql_num(unit['very_difficult_mh'])}, {sql_str(unit['verification_status'])}, "
            f"{'true' if unit['production_allowed'] else 'false'}, {sql_str(unit['notes'])})"
        )

    SQL.write_text(
        "-- Full Labor Database from Estim8r_Source_Audited_Electrical_Labor_Database.xlsx (1,921 rows).\n"
        "-- Completes the 5-row production-ready seed with the remaining model-baseline taxonomy.\n\n"
        "insert into public.labor_items (\n"
        "  id, trade, category, subcategory, item_name, description, material_type, size, unit, quantity_per_labor_unit, default_crew, active\n"
        ") values\n"
        + ",\n".join(item_values)
        + "\non conflict (id) do nothing;\n\n"
        "insert into public.labor_units (\n"
        "  id, labor_item_id, source_type, source_name, source_year, source_reference,\n"
        "  normal_mh, difficult_mh, very_difficult_mh, verification_status, production_allowed, notes\n"
        ") values\n"
        + ",\n".join(unit_values)
        + "\non conflict (id) do nothing;\n",
        encoding="utf-8",
    )
    print(f"wrote {len(items)} items from {path}")
    print(f"  {DATA_JS}")
    print(f"  {SQL}")


if __name__ == "__main__":
    main()
